"""
Export service — generate DOCX and PDF files from plain-text drafts.
"""

from __future__ import annotations

import io
import re
import textwrap


def export_docx(content: str, title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=0)

    for paragraph in content.split("\n\n"):
        para = paragraph.strip()
        if para:
            p = doc.add_paragraph(para)
            p.style.font.size = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_pdf(content: str, title: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", style="B", size=16)
    pdf.multi_cell(0, 10, title)
    pdf.ln(4)
    pdf.set_font("Helvetica", size=12)

    for paragraph in content.split("\n\n"):
        para = paragraph.strip()
        if para:
            # FPDF2 multi_cell handles line-wrapping automatically
            pdf.multi_cell(0, 8, para)
            pdf.ln(4)

    return bytes(pdf.output())
